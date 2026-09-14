# -*- coding: utf-8 -*-
import os, sys
import docx
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_val in borders.items():
        if border_val:
            b_xml = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="{border_val.get("val", "single")}" w:sz="{border_val.get("sz", "4")}" w:space="0" w:color="{border_val.get("color", "CCCCCC")}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def add_styled_run(paragraph, text, font_name="Bpmf Iansui Regular", font_size=12, bold=False, color_rgb=(40,40,40)):
    r = paragraph.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*color_rgb)
    if r._element.rPr is None:
        r._element.get_or_add_rPr()
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    r._element.rPr.rFonts.set(qn('w:ascii'), font_name)
    r._element.rPr.rFonts.set(qn('w:hAnsi'), font_name)
    return r

def create_worksheet(output_docx, is_teacher=False, img_path=""):
    doc = docx.Document()
    
    # 1. Page Setup - A4 Portrait with compact margins
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)
    
    # Header Paragraph
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    badge = "【教師/家長示範解答版】" if is_teacher else "【語文素養寫作篇】"
    add_styled_run(title_p, f"國小看圖寫話與好句子魔法學習單 {badge}", font_name="微軟正黑體", font_size=16, bold=True, color_rgb=(20, 70, 120))
    
    # Student Info Table
    info_table = doc.add_table(rows=1, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    cell_info = info_table.cell(0, 0)
    cell_info.width = Cm(13.0)
    p_info = cell_info.paragraphs[0]
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(0)
    add_styled_run(p_info, "二年____班   座號：____   姓名：____________", font_name="Bpmf Iansui Regular", font_size=11, bold=True, color_rgb=(60, 60, 60))
    
    cell_grade = info_table.cell(0, 1)
    cell_grade.width = Cm(5.4)
    p_grade = cell_grade.paragraphs[0]
    p_grade.paragraph_format.space_before = Pt(0)
    p_grade.paragraph_format.space_after = Pt(0)
    p_grade.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_styled_run(p_grade, "評語與星級：⭐⭐⭐⭐⭐", font_name="微軟正黑體", font_size=10, bold=False, color_rgb=(100, 100, 100))
    
    set_cell_borders(cell_info)
    set_cell_borders(cell_grade)

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(2)
    div_p.paragraph_format.space_after = Pt(4)
    div_run = add_styled_run(div_p, "━" * 46, font_name="微軟正黑體", font_size=8, bold=False, color_rgb=(180, 200, 220))
    div_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # SECTION 1: Image & Observation Helper (2-Column Table)
    top_table = doc.add_table(rows=1, cols=2)
    top_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    top_table.autofit = False
    
    # Left Cell: Image
    cell_img = top_table.cell(0, 0)
    cell_img.width = Cm(8.8)
    set_cell_background(cell_img, "FAFAFA")
    set_cell_margins(cell_img, top=60, bottom=60, left=60, right=60)
    set_cell_borders(cell_img, 
                     top={'val': 'single', 'sz': '4', 'color': 'D0D8E0'},
                     bottom={'val': 'single', 'sz': '4', 'color': 'D0D8E0'},
                     left={'val': 'single', 'sz': '4', 'color': 'D0D8E0'},
                     right={'val': 'single', 'sz': '4', 'color': 'D0D8E0'})
    
    p_img = cell_img.paragraphs[0]
    p_img.paragraph_format.space_before = Pt(0)
    p_img.paragraph_format.space_after = Pt(0)
    p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(img_path):
        p_img.add_run().add_picture(img_path, width=Cm(8.4))
    
    # Right Cell: Eagle Eye Observation Guide
    cell_guide = top_table.cell(0, 1)
    cell_guide.width = Cm(9.6)
    set_cell_background(cell_guide, "F5F9FD")
    set_cell_margins(cell_guide, top=80, bottom=80, left=120, right=80)
    set_cell_borders(cell_guide, 
                     top={'val': 'single', 'sz': '4', 'color': 'BBD5EA'},
                     bottom={'val': 'single', 'sz': '4', 'color': 'BBD5EA'},
                     left={'val': 'single', 'sz': '4', 'color': 'BBD5EA'},
                     right={'val': 'single', 'sz': '4', 'color': 'BBD5EA'})
    
    p_guide_title = cell_guide.paragraphs[0]
    p_guide_title.paragraph_format.space_before = Pt(0)
    p_guide_title.paragraph_format.space_after = Pt(3)
    add_styled_run(p_guide_title, "🔍【小鷹眼看一看】情境引導", font_name="微軟正黑體", font_size=11, bold=True, color_rgb=(25, 85, 145))
    
    observations = [
        ("🌸 時間與地點：", "溫暖的春日午後，美麗的公園草地。"),
        ("🪁 男孩在做什麼：", "笑瞇瞇地迎著微風奔跑，放著高高的風箏。"),
        ("📖 女孩與小狗：", "坐在樹下專心看故事書，金黃小狗乖乖陪伴。"),
        ("🦋 周圍還有：", "五顏六色的花朵盛開，美麗的蝴蝶飛舞。")
    ]
    for label, desc in observations:
        p_item = cell_guide.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = Pt(14)
        add_styled_run(p_item, label, font_name="Bpmf Iansui Regular", font_size=10, bold=True, color_rgb=(40, 40, 40))
        add_styled_run(p_item, desc, font_name="Bpmf Iansui Regular", font_size=9.5, bold=False, color_rgb=(70, 70, 70))

    # SECTION 2: Writing Scaffold (Filling the blanks)
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(5)
    p_sec2.paragraph_format.space_after = Pt(2)
    add_styled_run(p_sec2, "🪄【魔法任務：句子拉拉鍊】看圖填一填，寫出好句子", font_name="微軟正黑體", font_size=12, bold=True, color_rgb=(180, 70, 20))
    
    # Prompt Instruction
    p_prompt = doc.add_paragraph()
    p_prompt.paragraph_format.space_before = Pt(0)
    p_prompt.paragraph_format.space_after = Pt(3)
    add_styled_run(p_prompt, "請仔細看圖並參考小鷹眼的提示，把下列句子補充完整，讓句子變生動長大喔！", font_name="Bpmf Iansui Regular", font_size=10, bold=False, color_rgb=(80, 80, 80))
    
    # Sentence Scaffold Box (Single big styled table)
    scaffold_table = doc.add_table(rows=4, cols=1)
    scaffold_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    scaffold_table.autofit = False
    
    scaffold_items = [
        ("① 【時間與天氣】", 
         "在一個\U000e01e1", 
         "陽光溫暖、微風徐徐" if is_teacher else "________________________", 
         "（天氣）的春天下午，"),
        
        ("② 【地點與主角】", 
         "小天和小安在", 
         "開滿五顏六色小花的大公園草地上" if is_teacher else "____________________________________", 
         "（美麗的地點）快樂玩耍。"),
        
        ("③ 【動作描寫】", 
         "活潑的小天在草地上", 
         "開心地迎著微風奔跑放風箏" if is_teacher else "________________________________", 
         "；\r安靜的小安則靠著大樹，", 
         "專心地翻看精彩的圖畫故事書" if is_teacher else "________________________________", 
         "。"),
        
        ("④ 【心情與總結】", 
         "可愛的金黃色小狗乖乖趴在身邊，大家都覺得", 
         "既充實又無比幸福" if is_teacher else "________________________", 
         "（心情如何）！")
    ]
    
    for row_idx, item in enumerate(scaffold_items):
        cell = scaffold_table.cell(row_idx, 0)
        cell.width = Cm(18.4)
        set_cell_background(cell, "FCFDFD" if row_idx % 2 == 0 else "F7FAFC")
        set_cell_margins(cell, top=70, bottom=70, left=100, right=80)
        set_cell_borders(cell, 
                         bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                         left={'val': 'single', 'sz': '12', 'color': '4A90E2'})
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = Pt(18)
        
        # Title badge
        add_styled_run(p, item[0] + " ", font_name="Bpmf Iansui Regular", font_size=10.5, bold=True, color_rgb=(20, 90, 150))
        
        if len(item) == 4:
            add_styled_run(p, item[1], font_name="Bpmf Iansui Regular", font_size=11, bold=False, color_rgb=(30, 30, 30))
            # Blank or teacher answer
            ans_color = (180, 40, 20) if is_teacher else (0, 80, 180)
            add_styled_run(p, item[2], font_name="Bpmf Iansui Regular", font_size=11, bold=is_teacher, color_rgb=ans_color)
            add_styled_run(p, item[3], font_name="Bpmf Iansui Regular", font_size=11, bold=False, color_rgb=(30, 30, 30))
        elif len(item) == 6:
            add_styled_run(p, item[1], font_name="Bpmf Iansui Regular", font_size=11, bold=False, color_rgb=(30, 30, 30))
            ans_color = (180, 40, 20) if is_teacher else (0, 80, 180)
            add_styled_run(p, item[2], font_name="Bpmf Iansui Regular", font_size=11, bold=is_teacher, color_rgb=ans_color)
            add_styled_run(p, item[3], font_name="Bpmf Iansui Regular", font_size=11, bold=False, color_rgb=(30, 30, 30))
            add_styled_run(p, item[4], font_name="Bpmf Iansui Regular", font_size=11, bold=is_teacher, color_rgb=ans_color)
            add_styled_run(p, item[5], font_name="Bpmf Iansui Regular", font_size=11, bold=False, color_rgb=(30, 30, 30))

    # SECTION 3: Bottom Check & Tips (2 Columns Table)
    doc.add_paragraph().paragraph_format.space_before = Pt(2)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    
    bot_table = doc.add_table(rows=1, cols=2)
    bot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bot_table.autofit = False
    
    # Left: Self Assessment (Checklist)
    cell_eval = bot_table.cell(0, 0)
    cell_eval.width = Cm(10.2)
    set_cell_background(cell_eval, "FFFAFA")
    set_cell_margins(cell_eval, top=70, bottom=70, left=90, right=70)
    set_cell_borders(cell_eval, 
                     top={'val': 'single', 'sz': '4', 'color': 'F5D0D0'},
                     bottom={'val': 'single', 'sz': '4', 'color': 'F5D0D0'},
                     left={'val': 'single', 'sz': '4', 'color': 'F5D0D0'},
                     right={'val': 'single', 'sz': '4', 'color': 'F5D0D0'})
    
    p_eval_title = cell_eval.paragraphs[0]
    p_eval_title.paragraph_format.space_before = Pt(0)
    p_eval_title.paragraph_format.space_after = Pt(2)
    add_styled_run(p_eval_title, "⭐【我是自評小達人】寫完檢查打勾勾", font_name="微軟正黑體", font_size=10.5, bold=True, color_rgb=(180, 50, 40))
    
    checks = [
        "□ 我有寫出完整\U000e01e1的時間、地點與人物做什麼。",
        "□ 我用上了形容詞（如：五顏六色、溫暖、活潑）。",
        "□ 我的句子通順，並正確加上了標點符號。",
        "□ 我的國字和注音寫得端正乾淨。"
    ]
    for chk in checks:
        p_c = cell_eval.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1)
        p_c.paragraph_format.line_spacing = Pt(13)
        add_styled_run(p_c, chk, font_name="Bpmf Iansui Regular", font_size=9, bold=False, color_rgb=(60, 60, 60))
        
    # Right: Parent / Teacher Guiding Tips
    cell_tips = bot_table.cell(0, 1)
    cell_tips.width = Cm(8.2)
    set_cell_background(cell_tips, "F6FAF5")
    set_cell_margins(cell_tips, top=70, bottom=70, left=90, right=70)
    set_cell_borders(cell_tips, 
                     top={'val': 'single', 'sz': '4', 'color': 'D0E8D0'},
                     bottom={'val': 'single', 'sz': '4', 'color': 'D0E8D0'},
                     left={'val': 'single', 'sz': '4', 'color': 'D0E8D0'},
                     right={'val': 'single', 'sz': '4', 'color': 'D0E8D0'})
    
    p_tips_title = cell_tips.paragraphs[0]
    p_tips_title.paragraph_format.space_before = Pt(0)
    p_tips_title.paragraph_format.space_after = Pt(2)
    add_styled_run(p_tips_title, "💡【親師引導小錦囊】孩子卡關時問這三句：", font_name="微軟正黑體", font_size=10.5, bold=True, color_rgb=(35, 115, 60))
    
    tips = [
        "1. 「看圖裡的小男孩，他的風箏飛在什麼顏色的天上？」",
        "2. 「小女孩在看書，身邊的小狗表情看起來怎麼樣？」",
        "3. 「如果今天是你去這個公園，你的心裡會感覺如何？」"
    ]
    for tip in tips:
        p_t = cell_tips.add_paragraph()
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(1)
        p_t.paragraph_format.line_spacing = Pt(13)
        add_styled_run(p_t, tip, font_name="Bpmf Iansui Regular", font_size=8.5, bold=False, color_rgb=(50, 70, 50))
        
    # Bottom encourage footer
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(3)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_styled_run(p_foot, "✨ 寫作小叮嚀：寫完請自己大聲朗讀一遍，聽聽看句子是不是順暢好聽喔！ ✨", font_name="Bpmf Iansui Regular", font_size=9, bold=True, color_rgb=(120, 100, 30))

    doc.save(output_docx)
    print(f"Saved: {output_docx}")

if __name__ == '__main__':
    gdrive_dir = r"G:\我的雲端硬碟\低年級寫作學習單"
    img_path = os.path.join(gdrive_dir, "writing_scene.jpg")
    
    student_docx = os.path.join(gdrive_dir, "看圖寫話學習單_學生版.docx")
    teacher_docx = os.path.join(gdrive_dir, "看圖寫話學習單_教師版.docx")
    
    create_worksheet(student_docx, is_teacher=False, img_path=img_path)
    create_worksheet(teacher_docx, is_teacher=True, img_path=img_path)
